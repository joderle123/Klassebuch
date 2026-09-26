#!/usr/bin/env python3
"""Erzeugt import-test.xlsx für hub-quellen/tests/datenbank.js (Import aus einer Excel-Tabelle).
Nur erfundene Personen. Aufruf: python3 import-test.py  (braucht openpyxl)
Absichtlich „wie im Alltag“: Titelzeile über der Kopfzeile, französische Überschriften,
Name und Vorname in einer Spalte, echte Excel-Datumszellen und Datum als Text, DR als Zahl,
mehrere Maßnahmen in einer Zelle, eine unbekannte Spalte, ein unmögliches Datum, eine Zeile ohne Namen."""
import os, datetime as dt
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = 'Élèves 2025-26'
ws['A1'] = 'Liste CDSE 2025-26 (fictive – données de test)'
kopf = ['Nom et prénom', 'Date de naissance', 'Sexe', 'Matricule', 'École', 'Classe', 'DR', 'Mesures', 'Début ISA',
        'Atelier', 'Début atelier', 'Groupe CST', 'CC', 'QI', 'ELDiB', 'ELDiB V', 'ELDiB K', 'Remarques', 'Stelle', 'Couleur préférée']
ws.append([])
ws.append(kopf)
D = lambda y, m, d: dt.datetime(y, m, d)
zeilen = [
    ['MUSTER Tom', D(2014, 3, 10), 'garçon', '2014031000012', 'École Brill', 'C4.1', 6, 'ISA; Atelier', D(2025, 10, 15),
     'Atelier Lecture (fictif)', D(2025, 11, 3), None, 'CL', 95, D(2025, 12, 1), 3, 2, 'fiktiver Testfall', 'ISA', 'bleu'],
    ['BEISPIEL Nina', D(2016, 6, 15), 'fille', '2016061500099', 'École Test Sanem (fictive)', 'C3.2', 'DR 05', 'C&G parents, Rééducation', None,
     None, None, None, 'CDI; Centre de logopédie', None, D(2026, 3, 2), 2, 2, None, None, 'vert'],
    ['Probe, Jana', '12.02.2015', 'F', '', 'École Test Mersch (fictive)', 'C4.2', 'Mersch', 'CST', None,
     None, None, 'Passo', None, 105, None, None, None, None, None, 'rouge'],
    ['Fiktiv Max', '31.02.2015', 'm', '', 'École Test (fictive)', 'C2.1', '', 'DS', None,
     None, None, None, None, None, None, None, None, None, None, None],
    [None, None, None, None, None, 'C1.1', None, None, None, None, None, None, None, None, None, None, None, 'Zeile ohne Namen', None, None],
]
for z in zeilen:
    ws.append(z)
for zeile in ws.iter_rows(min_row=4):
    for c in zeile:
        if isinstance(c.value, dt.datetime):
            c.number_format = 'DD.MM.YYYY'
ws2 = wb.create_sheet('Notes')
ws2['A1'] = 'Notizen'
ws2['A2'] = 'nur Text, keine Schüler'
ziel = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'import-test.xlsx')
wb.save(ziel)
print('geschrieben:', ziel)
